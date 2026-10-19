#!/usr/bin/env python3
"""Read-only monthly accounting-actuals extractor for the TPO FS workbook.

Reads the ``PL``, ``BS`` and ``IncomeStatement`` tabs of the accounting export
without modifying it and produces:

* ``<prefix>.json`` — a narrow ``tpo-monthly-actuals-v1`` import pack with
  integer-satang amounts and full per-cell source references;
* ``<prefix>.xlsx`` — a three-sheet human review copy (MonthlyFinancials,
  1. Working Capital, Source Map).

The source workbook is never written, recalculated or macro-executed, and the
existing compensation-model outputs are never treated as actuals.  Public API:
``extract_monthly_inputs(source) -> dict``, ``write_review_workbook(pack,
output) -> None`` and ``main(argv) -> int``.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import sys
import uuid
import warnings
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

warnings.filterwarnings('ignore', message='Data Validation extension is not supported')

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.formula import ArrayFormula
except ImportError as exc:  # pragma: no cover - environment guard
    raise SystemExit(
        'openpyxl is required. Install tools/requirements-monthly-input.txt into an '
        'isolated environment, never system pip, e.g.:\n'
        '  python3 -m venv /tmp/tpo-monthly-import-venv\n'
        '  /tmp/tpo-monthly-import-venv/bin/python -m pip install -r '
        'tools/requirements-monthly-input.txt') from exc

SCHEMA = 'tpo-monthly-actuals-v1'
MAX_MONTHS = 24
CENT = Decimal('0.01')
TOLERANCE_CENTS = 1  # one satang

PL_SHEET, BS_SHEET, IS_SHEET = 'PL', 'BS', 'IncomeStatement'
TITLE_MONTHS = ['january', 'february', 'march', 'april', 'may', 'june', 'july',
                'august', 'september', 'october', 'november', 'december']
DISPLAY_MONTHS = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep',
                  'Oct', 'Nov', 'Dec']

# Verified monthly blocks: current-year PL R:AC / BS D:O, prior-year PL AG:AR / BS Q:AB.
PL_CURRENT_COL, PL_PRIOR_COL, BS_CURRENT_COL, BS_PRIOR_COL = 18, 33, 4, 17
BLOCK_WIDTH = 12

PL_FIELDS = {'revenueCents': 15, 'cogsCents': 63, 'sgaCents': 170, 'netIncomeCents': 197}
PL_CHECKS = {'grossProfitCents': 65, 'ebitCents': 172}
SENTINEL_ROW = 107
PL_ANCHORS = {15: 'T-R002', 63: 'T-C', 65: 'T-N001', 107: '53011.09', 170: 'T-E',
              172: 'T-N002', 197: 'PL'}
PL_LABELS = {
    15: 'Total Revenues (Sales)',
    63: 'Total Cost of Goods Sold',
    65: 'Gross Profit',
    107: 'Salary - Management',
    170: 'Total Selling, General and Administrative Expenses',
    172: 'Profit (Loss) before Other incomes and expenses',
    197: 'Net Income',
}
BS_ANCHORS = {
    5: '11111', 6: '11112', 7: '11121.01', 8: '11122.01', 9: '11311',
    12: '11511', 13: '11521', 14: '11541',
    48: '21311', 49: '21399', 50: '21911.99', 51: '21919', 52: '21921',
    53: '21941', 54: '21942', 55: '21951.01', 56: '21951.03', 57: '21951.99',
    58: '21952.01', 59: '21953', 60: '27111', 61: '27113', 62: '27121',
    63: '27123', 64: '27124', 65: '27125', 66: '27140',
}
BS_TOTAL_ROW, BS_BORROWING_ROWS = 68, (53, 54)
CASH_ROWS, AR_ROW, INVENTORY_ROWS = (5, 6, 7, 8), 9, (12, 13, 14)
AP_ROWS = tuple(range(48, 53)) + tuple(range(55, 67))  # 48–52 and 55–66

PL_TITLE_RE = re.compile(r'^profit\s*&\s*\(?\s*loss\s*\)?\s*report\s*:\s*ytd\s+([a-z]+)\s+(\d{4})$')
BS_TITLE_RE = re.compile(r'^balance\s*sheet\s*:\s*ytd\s+([a-z]+)\s+(\d{4})$')

REVIEW_NOTE = ('Review copy only. Import the JSON pack; do not replace the report '
               'workbook or paste over calculated columns.')
MONEY_FORMAT = '#,##0.00;[Red](#,##0.00);0.00'


class MonthlyInputError(Exception):
    """Validation failure; the message pinpoints the offending source cell(s)."""


def _norm_text(value) -> str:
    """Collapse all whitespace (including NBSP) for tolerant comparisons."""
    return re.sub(r'\s+', ' ', str(value or '').replace(' ', ' ')).strip()


def _norm_label(value) -> str:
    """Case/whitespace/comma-insensitive account label key."""
    return re.sub(r'[\s,]+', '', _norm_text(value).lower())


def _norm_code(value) -> str:
    """Account-code key that preserves numeric decimal identity."""
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, (int, float)):
        if not math.isfinite(value):
            return 'NaN'
        return format(Decimal(str(value)).normalize(), 'f')
    return _norm_text(value).upper()


def _is_formula(value) -> bool:
    return isinstance(value, ArrayFormula) or (
        isinstance(value, str) and value.startswith('='))


def _sheet_value(view, sheet: str, ref: str):
    return view[sheet][ref].value


def _require_sheets(wb, source: Path):
    names = set(wb.sheetnames)
    for name in (PL_SHEET, BS_SHEET, IS_SHEET):
        if name not in names:
            raise MonthlyInputError(
                f'{source.name}: required sheet "{name}" is missing. The supported '
                'source is the inspected PL/BS/IncomeStatement accounting export.')




def _read_title(view, sheet: str, ref: str, pattern: re.Pattern):
    raw = _sheet_value(view, sheet, ref)
    text = _norm_text(raw).lower()
    match = pattern.match(text)
    if not match:
        raise MonthlyInputError(
            f'{sheet}!{ref}: unrecognized title {raw!r}. Expected the '
            '"... : YTD <English month> <YYYY>" pattern of the supported export.')
    month_name, year_text = match.groups()
    if month_name not in TITLE_MONTHS:
        raise MonthlyInputError(
            f'{sheet}!{ref}: unknown English month {month_name!r} in the report title.')
    return TITLE_MONTHS.index(month_name) + 1, int(year_text)



def _check_year_anchors(formula_view, current_year: int):
    prev_year = current_year - 1
    cur_yy, prev_yy = f'{current_year % 100:02d}', f'{prev_year % 100:02d}'
    checks = [
        (PL_SHEET, 'AD3', f'YTD-{cur_yy}', 'current-year YTD anchor'),
        (PL_SHEET, 'AS3', f'YTD-{prev_yy}', 'prior-year YTD anchor'),
    ]
    for sheet, ref, expected, label in checks:
        actual = _norm_text(_sheet_value(formula_view, sheet, ref))
        if actual != expected:
            raise MonthlyInputError(
                f'{sheet}!{ref}: {label} is {actual!r}; expected {expected!r}. '
                'A shifted or unsupported layout stops extraction.')
    for sheet, ref, expected_year in ((IS_SHEET, 'R6', current_year),
                                      (IS_SHEET, 'AS6', prev_year)):
        raw = _sheet_value(formula_view, sheet, ref)
        year = None
        if isinstance(raw, bool):
            year = None
        elif isinstance(raw, int):
            year = raw
        elif isinstance(raw, float) and math.isfinite(raw) and raw.is_integer():
            year = int(raw)
        elif isinstance(raw, str) and raw.strip().isdigit():
            year = int(raw.strip())
        if year != expected_year:
            raise MonthlyInputError(
                f'{sheet}!{ref}: year anchor is {raw!r}; expected {expected_year}. '
                'The stored header years are not authoritative; these anchors are.')
    for sheet, ref, expected in ((IS_SHEET, 'R7', f'jan-{cur_yy}'),
                                 (IS_SHEET, 'AS7', f'jan-{prev_yy}')):
        actual = _norm_text(_sheet_value(formula_view, sheet, ref)).lower()
        if actual != expected:
            raise MonthlyInputError(
                f'{sheet}!{ref}: month anchor is {actual!r}; expected {expected!r}.')


def _check_block_headers(formula_view, sheet: str, start_col: int, year: int):
    """Verify January–December order using only each header's month component."""
    for offset in range(BLOCK_WIDTH):
        col = get_column_letter(start_col + offset)
        ref = f'{col}3'
        value = _sheet_value(formula_view, sheet, ref)
        if not isinstance(value, datetime):
            raise MonthlyInputError(
                f'{sheet}!{ref}: expected a date-typed month header, found '
                f'{value!r}. The stored year is never used; the month component '
                'must identify January–December in order.')
        if value.month != offset + 1:
            raise MonthlyInputError(
                f'{sheet}!{ref}: month component is {value.month}; expected '
                f'{offset + 1} for this verified {year} block.')


def _classify_sentinel(formula_view, cached_view, sheet: str, start_col: int,
                       year: int) -> int:
    """Return the number of reported months (contiguous from January)."""
    reported = []
    for offset in range(BLOCK_WIDTH):
        col = get_column_letter(start_col + offset)
        coord = f'{col}{SENTINEL_ROW}'
        ref = f'{sheet}!{coord}'
        formula = _sheet_value(formula_view, sheet, coord)
        cached = _sheet_value(cached_view, sheet, coord)
        if _is_formula(formula):
            if isinstance(cached, bool) or not isinstance(cached, (int, float)) \
                    or not math.isfinite(cached):
                raise MonthlyInputError(
                    f'{ref}: salary-sentinel formula has no numeric cached value. '
                    'A formula sentinel is an error, not an unreported month; '
                    'recalculate and save the source in Excel before extraction.')
            reported.append(offset + 1)
        elif cached is None:
            continue  # non-formula blank: unreported month
        elif isinstance(cached, (int, float)) and not isinstance(cached, bool) \
                and math.isfinite(cached):
            reported.append(offset + 1)  # numeric literal, genuine zero included
        else:
            raise MonthlyInputError(
                f'{ref}: unsupported salary sentinel {cached!r}.')
    expected_prefix = list(range(1, len(reported) + 1))
    if reported != expected_prefix:
        first_gap = next((reported[i] for i in range(len(reported))
                          if reported[i] != i + 1), reported[-1] if reported else 1)
        col = get_column_letter(start_col + first_gap - 1)
        raise MonthlyInputError(
            f'{sheet}!{col}{SENTINEL_ROW}: populated salary sentinel after a gap '
            f'in the {year} block. Actual months must be contiguous from January.')
    return len(reported)


def _check_pl_anchors(formula_view):
    for row, code in PL_ANCHORS.items():
        actual = _norm_code(_sheet_value(formula_view, PL_SHEET, f'A{row}'))
        if actual != code:
            raise MonthlyInputError(
                f'{PL_SHEET}!A{row}: account anchor is {actual!r}; expected '
                f'{code!r}. A shifted or unsupported layout stops extraction.')
        label = _norm_label(_sheet_value(formula_view, PL_SHEET, f'B{row}'))
        if label != _norm_label(PL_LABELS[row]):
            raise MonthlyInputError(
                f'{PL_SHEET}!B{row}: anchor label does not match the supported '
                f'layout (expected "{PL_LABELS[row]}").')


def _check_bs_anchors(formula_view):
    for row, code in BS_ANCHORS.items():
        actual = _norm_code(_sheet_value(formula_view, BS_SHEET, f'A{row}'))
        if actual != code:
            raise MonthlyInputError(
                f'{BS_SHEET}!A{row}: account anchor is {actual!r}; expected '
                f'{code!r}. A shifted or unsupported layout stops extraction.')
    if _norm_label(_sheet_value(formula_view, BS_SHEET, 'B68')) != 'totalliabilities':
        raise MonthlyInputError(
            f'{BS_SHEET}!B68: expected the "Total Liabilities" subtotal label.')
    for row in BS_BORROWING_ROWS:
        label = _norm_text(_sheet_value(formula_view, BS_SHEET, f'B{row}')).lower()
        if 'borrowing' not in label:
            raise MonthlyInputError(
                f'{BS_SHEET}!B{row}: expected a borrowing label so it can be '
                'excluded from accounts payable.')
    ar_label = _norm_text(_sheet_value(formula_view, BS_SHEET, 'B9')).lower()
    if 'trade' not in ar_label or 'receivable' not in ar_label:
        raise MonthlyInputError(
            f'{BS_SHEET}!B9: expected a trade-receivables label.')
    for row, words in ((12, ('finished', 'goods')), (13, ('raw', 'material')),
                       (14, ('product', 'packing', 'material'))):
        label = _norm_text(_sheet_value(formula_view, BS_SHEET, f'B{row}')).lower()
        if not all(word in label for word in words):
            raise MonthlyInputError(
                f'{BS_SHEET}!B{row}: inventory-component label does not match the '
                'supported layout.')


def _amount(cached_view, formula_view, sheet: str, row: int, col: int,
            blank_zero: bool = False) -> Decimal | None:
    """Read one accounting amount in THB as an exact Decimal.

    Booleans, text, non-finite numbers and formulas without numeric caches are
    errors.  Only AP detail rows accept a genuinely absent (non-formula) blank
    as a zero contribution; everywhere else a blank is an error, not zero.
    """
    coord = f'{get_column_letter(col)}{row}'
    ref = f'{sheet}!{coord}'
    formula = _sheet_value(formula_view, sheet, coord)
    cached = _sheet_value(cached_view, sheet, coord)
    if isinstance(formula, bool) or isinstance(cached, bool):
        raise MonthlyInputError(f'{ref}: boolean value is not a valid amount.')
    if cached is None:
        if _is_formula(formula):
            raise MonthlyInputError(
                f'{ref}: formula has no numeric cached value. Recalculate and save '
                'the source workbook in Excel before extraction.')
        if blank_zero:
            return Decimal('0')
        raise MonthlyInputError(f'{ref}: required amount is blank; blank is not zero.')
    if isinstance(cached, str):
        raise MonthlyInputError(
            f'{ref}: cached value {cached!r} is not numeric (formula error or text).')
    if not isinstance(cached, (int, float)) or not math.isfinite(cached):
        raise MonthlyInputError(f'{ref}: non-finite numeric value {cached!r}.')
    return Decimal(str(cached))


def _to_cents(amounts) -> int:
    """Sum Decimals in THB, then round the field total once, HALF_UP, to satang."""
    total = sum(amounts, Decimal('0'))
    return int((total.quantize(CENT, rounding=ROUND_HALF_UP) * 100).to_integral_value())


def _month_key(year: int, month: int) -> str:
    return f'{year:04d}-{month:02d}'


def _month_label(key: str) -> str:
    year, month = key.split('-')
    return f'{DISPLAY_MONTHS[int(month) - 1]}-{year[2:]}'


def _quarter(key: str) -> str:
    year, month = key.split('-')
    return f'Q{(int(month) + 2) // 3} {year}'


def extract_monthly_inputs(source: Path) -> dict:
    """Validate the accounting export read-only and build the import pack."""
    source = Path(source)
    if not source.is_file():
        raise MonthlyInputError(f'{source}: source workbook not found.')
    sha256 = hashlib.sha256(source.read_bytes()).hexdigest()
    formula_view = None
    cached_view = None
    try:
        formula_view = openpyxl.load_workbook(source, read_only=True, keep_links=False)
        cached_view = openpyxl.load_workbook(source, read_only=True, data_only=True,
                                             keep_links=False)
        _require_sheets(formula_view, source)
        pl_month, pl_year = _read_title(cached_view, PL_SHEET, 'A2', PL_TITLE_RE)
        bs_month, bs_year = _read_title(cached_view, BS_SHEET, 'A2', BS_TITLE_RE)
        if (pl_month, pl_year) != (bs_month, bs_year):
            raise MonthlyInputError(
                f'{PL_SHEET}!A2 / {BS_SHEET}!A2: report titles disagree '
                f'({pl_year}-{pl_month:02d} vs {bs_year}-{bs_month:02d}). '
                'A filename date is not evidence of a reporting period.')
        current_year = pl_year
        _check_year_anchors(formula_view, current_year)
        _check_block_headers(formula_view, PL_SHEET, PL_CURRENT_COL, current_year)
        _check_block_headers(formula_view, PL_SHEET, PL_PRIOR_COL, current_year - 1)
        _check_block_headers(formula_view, BS_SHEET, BS_CURRENT_COL, current_year)
        _check_block_headers(formula_view, BS_SHEET, BS_PRIOR_COL, current_year - 1)
        current_count = _classify_sentinel(formula_view, cached_view, PL_SHEET,
                                           PL_CURRENT_COL, current_year)
        prior_count = _classify_sentinel(formula_view, cached_view, PL_SHEET,
                                         PL_PRIOR_COL, current_year - 1)
        if current_count != pl_month:
            col = get_column_letter(PL_CURRENT_COL + max(current_count, pl_month - 1))
            raise MonthlyInputError(
                f'{PL_SHEET}!{col}{SENTINEL_ROW}: the current-year actual run ends at '
                f'month {current_count}, but the title reports {pl_month}. Later '
                'sentinels must be blank and the run must end exactly at the '
                'reporting month.')
        if prior_count != BLOCK_WIDTH:
            col = get_column_letter(PL_PRIOR_COL + prior_count)
            raise MonthlyInputError(
                f'{PL_SHEET}!{col}{SENTINEL_ROW}: prior comparison year '
                f'{current_year - 1} has only {prior_count}/12 reported months. '
                'Reject incomplete prior-year data rather than inventing coverage.')
        _check_pl_anchors(formula_view)
        _check_bs_anchors(formula_view)

        months = []
        for month in range(1, BLOCK_WIDTH + 1):
            months.append((_month_key(current_year - 1, month),
                           PL_PRIOR_COL + month - 1, BS_PRIOR_COL + month - 1))
        for month in range(1, current_count + 1):
            months.append((_month_key(current_year, month),
                           PL_CURRENT_COL + month - 1, BS_CURRENT_COL + month - 1))
        keys = [key for key, _, _ in months]
        if len(set(keys)) != len(keys) or len(keys) > MAX_MONTHS:
            raise MonthlyInputError(
                'Extracted month set is duplicated or exceeds 24 unique months.')

        records = []
        for key, pl_col, bs_col in months:
            months_ref = f'{key} ({_month_label(key)})'
            amounts = {}
            for field, row in PL_FIELDS.items():
                amounts[field] = [_amount(cached_view, formula_view, PL_SHEET, row, pl_col)]
            for field, row in PL_CHECKS.items():
                amounts[field] = [_amount(cached_view, formula_view, PL_SHEET, row, pl_col)]
            cash = [_amount(cached_view, formula_view, BS_SHEET, row, bs_col)
                    for row in CASH_ROWS]
            ar = [_amount(cached_view, formula_view, BS_SHEET, AR_ROW, bs_col)]
            inventory = [_amount(cached_view, formula_view, BS_SHEET, row, bs_col)
                         for row in INVENTORY_ROWS]
            ap = [_amount(cached_view, formula_view, BS_SHEET, row, bs_col,
                          blank_zero=True) for row in AP_ROWS]
            liabilities = [_amount(cached_view, formula_view, BS_SHEET, BS_TOTAL_ROW, bs_col)]
            borrowings = [_amount(cached_view, formula_view, BS_SHEET, row, bs_col)
                          for row in BS_BORROWING_ROWS]

            def cents(parts):
                return _to_cents(parts)

            record = {
                'month': key,
                'financials': {
                    'revenueCents': cents(amounts['revenueCents']),
                    'cogsCents': cents(amounts['cogsCents']),
                    'sgaCents': cents(amounts['sgaCents']),
                    'netIncomeCents': cents(amounts['netIncomeCents']),
                },
                'workingCapital': {
                    'cashCents': cents(cash),
                    'arCents': cents(ar),
                    'inventoryCents': cents(inventory),
                    'apCents': cents(ap),
                },
                'checks': {
                    'grossProfitCents': cents(amounts['grossProfitCents']),
                    'ebitCents': cents(amounts['ebitCents']),
                },
                'sourceRefs': {
                    'revenueCents': [f'{PL_SHEET}!{get_column_letter(pl_col)}{PL_FIELDS["revenueCents"]}'],
                    'cogsCents': [f'{PL_SHEET}!{get_column_letter(pl_col)}{PL_FIELDS["cogsCents"]}'],
                    'sgaCents': [f'{PL_SHEET}!{get_column_letter(pl_col)}{PL_FIELDS["sgaCents"]}'],
                    'netIncomeCents': [f'{PL_SHEET}!{get_column_letter(pl_col)}{PL_FIELDS["netIncomeCents"]}'],
                    'cashCents': [f'{BS_SHEET}!{get_column_letter(bs_col)}{row}' for row in CASH_ROWS],
                    'arCents': [f'{BS_SHEET}!{get_column_letter(bs_col)}{AR_ROW}'],
                    'inventoryCents': [f'{BS_SHEET}!{get_column_letter(bs_col)}{row}' for row in INVENTORY_ROWS],
                    'apCents': [f'{BS_SHEET}!{get_column_letter(bs_col)}{row}' for row in AP_ROWS],
                },
            }
            fin, wc, checks = record['financials'], record['workingCapital'], record['checks']
            expected_gp_dec = amounts['revenueCents'][0] - amounts['cogsCents'][0]
            if abs(amounts['grossProfitCents'][0] - expected_gp_dec) > Decimal('0.01'):
                raise MonthlyInputError(
                    f'{months_ref}: Gross Profit raw identity fails '
                    f"({amounts['grossProfitCents'][0]} vs revenue−COGS "
                    f"{expected_gp_dec}). See "
                    f"{PL_SHEET}!{get_column_letter(pl_col)}{PL_CHECKS['grossProfitCents']}.")
            expected_ebit_dec = amounts['grossProfitCents'][0] - amounts['sgaCents'][0]
            if abs(amounts['ebitCents'][0] - expected_ebit_dec) > Decimal('0.01'):
                raise MonthlyInputError(
                    f'{months_ref}: EBIT raw identity fails ({amounts["ebitCents"][0]} vs '
                    f"GP−SG&A {expected_ebit_dec}). See "
                    f"{PL_SHEET}!{get_column_letter(pl_col)}{PL_CHECKS['ebitCents']}.")
            if abs(checks['grossProfitCents'] - (fin['revenueCents'] - fin['cogsCents'])) > TOLERANCE_CENTS:
                raise MonthlyInputError(
                    f'{months_ref}: Gross Profit identity fails '
                    f"({checks['grossProfitCents']} vs revenue−COGS "
                    f"{fin['revenueCents'] - fin['cogsCents']}). See "
                    f"{PL_SHEET}!{get_column_letter(pl_col)}{PL_CHECKS['grossProfitCents']}.")
            if abs(checks['ebitCents'] - (checks['grossProfitCents'] - fin['sgaCents'])) > TOLERANCE_CENTS:
                raise MonthlyInputError(
                    f'{months_ref}: EBIT identity fails ({checks["ebitCents"]} vs '
                    f"GP−SG&A {checks['grossProfitCents'] - fin['sgaCents']}). See "
                    f"{PL_SHEET}!{get_column_letter(pl_col)}{PL_CHECKS['ebitCents']}.")
            ap_detail_dec = sum(ap, Decimal(0))
            expected_ap_dec = liabilities[0] - borrowings[0] - borrowings[1]
            if abs(ap_detail_dec - expected_ap_dec) > Decimal('0.01'):
                raise MonthlyInputError(
                    f'{months_ref}: AP detail rows do not reconcile with '
                    f'{BS_SHEET}!{get_column_letter(bs_col)}{BS_TOTAL_ROW} minus the '
                    f'two borrowing rows ({ap_detail_dec} vs {expected_ap_dec}); a missing '
                    'detail amount is hidden by the total.')
            records.append(record)

        pack_months = sorted(records, key=lambda m: m['month'])

        return {
            'schema': SCHEMA,
            'source': {'filename': source.name, 'sha256': sha256,
                       'reportMonth': _month_key(current_year, pl_month)},
            'months': pack_months,
        }
    finally:
        if formula_view is not None:
            formula_view.close()
        if cached_view is not None:
            cached_view.close()


def _text_cell(ws, ref: str, value: str):
    """Store metadata as explicit Excel text; a leading '=' never executes."""
    cell = ws[ref]
    cell.value = value
    if isinstance(value, str):
        cell.data_type = 's'
        cell.number_format = '@'


def _style_header(ws, row: int, width: int):
    from openpyxl.styles import Alignment, Font, PatternFill
    for col in range(1, width + 1):
        cell = ws.cell(row, col)
        cell.fill = PatternFill('solid', fgColor='0B1F3A')
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.freeze_panes = f'A{row + 1}'


def _money_cell(ws, row: int, col: int, cents: int):
    from openpyxl.styles import Alignment
    cell = ws.cell(row, col, float(Decimal(cents) / Decimal(100)))
    cell.number_format = MONEY_FORMAT
    cell.alignment = Alignment(horizontal='right')


def write_review_workbook(pack: dict, output: Path) -> None:
    """Write the three-sheet human review copy (never an import source)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    fin = wb.active
    fin.title = 'MonthlyFinancials'
    fin.append(['Month', 'Total Revenue', 'COGS', 'Gross Profit', 'SG&A', 'EBIT',
                'Net Income', 'Quarter'])
    _style_header(fin, 1, 8)
    wc = wb.create_sheet('1. Working Capital')
    wc.append(['Reporting Month', 'Cash Balance', 'Accounts Receivable',
               'Inventory Value', 'Accounts Payable', 'Net Working Capital'])
    _style_header(wc, 1, 6)
    for record in pack['months']:
        key, f, w, c = record['month'], record['financials'], record['workingCapital'], record['checks']
        row = fin.max_row + 1
        _text_cell(fin, f'A{row}', _month_label(key))
        gp = f['revenueCents'] - f['cogsCents']
        ebit = gp - f['sgaCents']
        for col, cents in enumerate((f['revenueCents'], f['cogsCents'],
                                     gp, f['sgaCents'],
                                     ebit, f['netIncomeCents']), start=2):
            _money_cell(fin, row, col, cents)
        _text_cell(fin, f'H{row}', _quarter(key))
        row = wc.max_row + 1
        _text_cell(wc, f'A{row}', _month_label(key))
        nwc = w['cashCents'] + w['arCents'] + w['inventoryCents'] - w['apCents']
        for col, cents in enumerate((w['cashCents'], w['arCents'], w['inventoryCents'],
                                     w['apCents'], nwc), start=2):
            _money_cell(wc, row, col, cents)
    for sheet, widths in ((fin, (12, 16, 14, 15, 14, 15, 15, 11)),
                          (wc, (15, 15, 19, 15, 17, 19))):
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width

    source = pack['source']
    sm = wb.create_sheet('Source Map')
    _text_cell(sm, 'A1', REVIEW_NOTE)
    sm['A1'].font = Font(bold=True, color='9C0006')
    sm['A1'].alignment = Alignment(wrap_text=True, vertical='top')
    sm.merge_cells('A1:D1')
    sm.row_dimensions[1].height = 30
    _text_cell(sm, 'A3', 'Source file')
    _text_cell(sm, 'B3', source['filename'])
    _text_cell(sm, 'A4', 'SHA-256')
    _text_cell(sm, 'B4', source['sha256'])
    _text_cell(sm, 'A5', 'Reporting month')
    _text_cell(sm, 'B5', source['reportMonth'])
    sm['A3'].font = sm['A4'].font = sm['A5'].font = Font(bold=True)
    sm.append([])
    header_row = 7
    for col, title in enumerate(('Month', 'Field', 'Source cells', 'Note'), start=1):
        sm.cell(header_row, col, title)
    _style_header(sm, header_row, 4)
    labels = {
        'revenueCents': 'Total Revenue', 'cogsCents': 'COGS', 'sgaCents': 'SG&A',
        'netIncomeCents': 'Net Income', 'cashCents': 'Cash Balance',
        'arCents': 'Accounts Receivable', 'inventoryCents': 'Inventory Value',
        'apCents': 'Accounts Payable',
    }
    row = header_row + 1
    for record in pack['months']:
        key = record['month']
        refs = record['sourceRefs']
        for field, label in labels.items():
            _text_cell(sm, f'A{row}', _month_label(key))
            _text_cell(sm, f'B{row}', label)
            _text_cell(sm, f'C{row}', ', '.join(refs[field]))
            _text_cell(sm, f'D{row}', '')
            row += 1
        bs_col = refs['arCents'][0].split('!')[1][:len(refs['arCents'][0].split('!')[1]) - len(str(AR_ROW))]
        _text_cell(sm, f'A{row}', _month_label(key))
        _text_cell(sm, f'B{row}', 'Borrowing excluded from Accounts Payable')
        _text_cell(sm, f'C{row}', ', '.join(
            f'{BS_SHEET}!{bs_col}{r}' for r in BS_BORROWING_ROWS))
        _text_cell(sm, f'D{row}', 'Excluded borrowing detail; not imported.')
        row += 1
        pl_col_letter = refs['revenueCents'][0].split('!')[1].rstrip('0123456789')
        for field, label in (('grossProfitCents', 'Gross Profit (identity check)'),
                             ('ebitCents', 'EBIT (identity check)')):
            _text_cell(sm, f'A{row}', _month_label(key))
            _text_cell(sm, f'B{row}', label)
            _text_cell(sm, f'C{row}', f'{PL_SHEET}!{pl_col_letter}{PL_CHECKS[field]}')
            _text_cell(sm, f'D{row}', 'Identity check — not imported.')
            row += 1
    for col, width in enumerate((12, 38, 52, 34), start=1):
        sm.column_dimensions[get_column_letter(col)].width = width
    wb.save(output)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(
        description='Extract a read-only monthly accounting-actuals JSON pack and '
                    'XLSX review copy from the TPO FS workbook.')
    parser.add_argument('--source', required=True,
                        help='accounting .xlsm/.xlsx export (never modified)')
    parser.add_argument('--output-prefix', required=True,
                        help='output path without extension; writes <prefix>.json '
                             'and <prefix>.xlsx')
    args = parser.parse_args(argv)
    source = Path(args.source)
    prefix = Path(args.output_prefix)
    outputs = [Path(f'{prefix}.json'), Path(f'{prefix}.xlsx')]
    try:
        if not source.is_file():
            raise MonthlyInputError(f'{source}: source workbook not found.')
        for out in outputs:
            if out.exists():
                raise MonthlyInputError(
                    f'{out}: output already exists; refusing to overwrite.')
        resolved_source = source.resolve()
        for out in outputs:
            if out.resolve(strict=False) == resolved_source:
                raise MonthlyInputError(
                    f'{out}: output must not resolve to the source workbook.')
        parent = prefix.parent if str(prefix.parent) else Path('.')
        if not parent.is_dir():
            raise MonthlyInputError(f'{parent}: output directory does not exist.')
        pack = extract_monthly_inputs(source)
        unique_id = f'{os.getpid()}_{uuid.uuid4().hex}'
        temp_paths = [parent / f'{prefix.name}.{unique_id}{suffix}.tmp' for suffix in ('.json', '.xlsx')]
        try:
            temp_paths[0].write_text(json.dumps(pack, indent=2) + '\n', encoding='utf-8')
            write_review_workbook(pack, temp_paths[1])
            os.replace(temp_paths[0], outputs[0])
            os.replace(temp_paths[1], outputs[1])
        except BaseException:
            for temp in temp_paths:
                try:
                    temp.unlink()
                except FileNotFoundError:
                    pass
            raise
    except MonthlyInputError as error:
        print(f'error: {error}', file=sys.stderr)
        return 2
    first, last = pack['months'][0]['month'], pack['months'][-1]['month']
    print(f"Extracted {len(pack['months'])} actual months ({first}..{last}) from "
          f"{pack['source']['filename']} (report {pack['source']['reportMonth']}).")
    print(f"Wrote {outputs[0]} and {outputs[1]}.")
    return 0


if __name__ == '__main__':
    sys.exit(main())
