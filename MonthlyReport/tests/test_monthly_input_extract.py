"""Behavior tests for tools/extract-monthly-input.py (standard library unittest).

Fixtures are synthetic workbooks built with the pinned openpyxl dependency in
temporary directories; the private accounting workbook is never used.  Cached
formula values and shared-formula followers are patched into the OpenXML zip
directly because openpyxl never calculates formulas.
"""

import importlib.util
import json
import re
import shutil
import sys
import tempfile
import unittest
import zipfile
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
_spec = importlib.util.spec_from_file_location(
    'extract_monthly_input', ROOT / 'tools' / 'extract-monthly-input.py')
extractor = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(extractor)

import openpyxl  # noqa: E402  (after the extractor's own import guard)

MONTH_NAMES = ['January', 'February', 'March', 'April', 'May', 'June', 'July',
               'August', 'September', 'October', 'November', 'December']
PL_ROWS = dict(extractor.PL_FIELDS, **extractor.PL_CHECKS)
AP_ROWS = extractor.AP_ROWS


def default_pl(month):
    rev = 1000 + month * 8 + 0.25
    cogs = 500 + month * 4 + 0.5
    sga = 400 + month * 2 + 0.75
    return {'revenueCents': rev, 'cogsCents': cogs, 'sgaCents': sga,
            'grossProfitCents': rev - cogs, 'ebitCents': rev - cogs - sga,
            'netIncomeCents': rev - cogs - sga - 10.25}


def default_bs(month):
    cash = [100.25, 200.5, 300.75, 400 + month]
    ar = 500 + month * 1.25
    inventory = [111.11 + month, 222.22 + month, 333.33 + month]
    ap = [10 + month * 0.01] * len(AP_ROWS)
    borrowing = [5000 + month, 2500]
    return {'cash': cash, 'ar': ar, 'inventory': inventory, 'ap': ap,
            'borrowing': borrowing}


def build_source(path, *, year=2027, report_month=3, pl=None, bs=None,
                 sentinels_current=None, sentinels_prior=None,
                 titles=None, overrides=None):
    """Create a synthetic accounting export matching the validated layout."""
    pl = pl or {}
    bs = bs or {}
    wb = openpyxl.Workbook()
    pl_ws = wb.active
    pl_ws.title = 'PL'
    bs_ws = wb.create_sheet('BS')
    is_ws = wb.create_sheet('IncomeStatement')
    prior_year = year - 1
    titles = titles or {}
    pl_ws['A2'] = titles.get('PL', f'Profit & (Loss) report  :  YTD '
                                   f'{MONTH_NAMES[report_month - 1]} {year}')
    bs_ws['A2'] = titles.get('BS', f'Balance Sheet  :  YTD '
                                   f'{MONTH_NAMES[report_month - 1]} {year}')
    pl_ws['AD3'] = titles.get('AD3', f'YTD-{year % 100:02d}')
    pl_ws['AS3'] = titles.get('AS3', f'YTD-{prior_year % 100:02d}')
    is_ws['R6'] = year
    is_ws['AS6'] = prior_year
    is_ws['R7'] = f'Jan-{year % 100:02d}'
    is_ws['AS7'] = f'Jan-{prior_year % 100:02d}'
    # Header dates deliberately carry a misleading stored year/day mix.
    for offset in range(12):
        pl_ws.cell(3, extractor.PL_CURRENT_COL + offset,
                   datetime(year + 1, offset + 1, 26 if offset % 2 else 25))
        pl_ws.cell(3, extractor.PL_PRIOR_COL + offset,
                   datetime(year, offset + 1, 25 if offset % 2 else 24))
        bs_ws.cell(3, extractor.BS_CURRENT_COL + offset,
                   datetime(year + 2, offset + 1, 26))
        bs_ws.cell(3, extractor.BS_PRIOR_COL + offset,
                   datetime(year, offset + 1, 25))
    for row, code in extractor.PL_ANCHORS.items():
        pl_ws.cell(row, 1, code if code != '53011.09' else 53011.09)
        pl_ws.cell(row, 2, extractor.PL_LABELS[row])
    for row, code in extractor.BS_ANCHORS.items():
        bs_ws.cell(row, 1, float(code) if '.' in code else int(code))
    bs_ws['B68'] = 'Total Liabilities'
    for row in extractor.BS_BORROWING_ROWS:
        bs_ws.cell(row, 2, f'Borrowing from party {row}')
    bs_ws['B9'] = 'Trade Receivables - General'
    bs_ws['B12'], bs_ws['B13'], bs_ws['B14'] = 'Finished Goods', 'Raw Material', \
        'Product packing materials'
    sentinels_current = sentinels_current if sentinels_current is not None else \
        [m <= report_month for m in range(1, 13)]
    sentinels_prior = sentinels_prior if sentinels_prior is not None else \
        [True] * 12
    for m in range(1, 13):
        pl_ws.cell(extractor.SENTINEL_ROW, extractor.PL_CURRENT_COL + m - 1,
                   100000 if sentinels_current[m - 1] else None)
        pl_ws.cell(extractor.SENTINEL_ROW, extractor.PL_PRIOR_COL + m - 1,
                   100000 if sentinels_prior[m - 1] else None)

    def put_month(m):
        values = pl.get(m, default_pl(m))
        balances = bs.get(m, default_bs(m))
        for col_base, key in ((extractor.PL_CURRENT_COL, 'cur'), (extractor.PL_PRIOR_COL, 'pri')):
            pl_ws.cell(PL_ROWS['revenueCents'], col_base + m - 1, values['revenueCents'])
            pl_ws.cell(PL_ROWS['cogsCents'], col_base + m - 1, values['cogsCents'])
            pl_ws.cell(PL_ROWS['sgaCents'], col_base + m - 1, values['sgaCents'])
            pl_ws.cell(PL_ROWS['grossProfitCents'], col_base + m - 1, values['grossProfitCents'])
            pl_ws.cell(PL_ROWS['ebitCents'], col_base + m - 1, values['ebitCents'])
            pl_ws.cell(PL_ROWS['netIncomeCents'], col_base + m - 1, values['netIncomeCents'])
            for i, row in enumerate(extractor.CASH_ROWS):
                bs_ws.cell(row, extractor.BS_CURRENT_COL + m - 1 if col_base == extractor.PL_CURRENT_COL
                           else extractor.BS_PRIOR_COL + m - 1, balances['cash'][i])
            wc_current = col_base == extractor.PL_CURRENT_COL
            bs_col = extractor.BS_CURRENT_COL + m - 1 if wc_current else extractor.BS_PRIOR_COL + m - 1
            bs_ws.cell(extractor.AR_ROW, bs_col, balances['ar'])
            for i, row in enumerate(extractor.INVENTORY_ROWS):
                bs_ws.cell(row, bs_col, balances['inventory'][i])
            total = sum(x for x in balances['ap'] if x is not None) + balances['borrowing'][0] + balances['borrowing'][1]
            for i, row in enumerate(AP_ROWS):
                bs_ws.cell(row, bs_col, balances['ap'][i])
            bs_ws.cell(53, bs_col, balances['borrowing'][0])
            bs_ws.cell(54, bs_col, balances['borrowing'][1])
            bs_ws.cell(extractor.BS_TOTAL_ROW, bs_col, total)

    for m in range(1, 13):
        put_month(m)
    for sheet, coord, value in (overrides or []):
        wb[sheet][coord] = value
    wb.save(path)
    return path


def sheet_xml_path(items, sheet):
    workbook = items['xl/workbook.xml'].decode('utf-8')
    match = re.search(r'<sheet[^>]*name="%s"[^>]*r:id="(rId\d+)"' % re.escape(sheet), workbook)
    if not match:
        match = re.search(r'<sheet[^>]*r:id="(rId\d+)"[^>]*name="%s"' % re.escape(sheet), workbook)
    rels = items['xl/_rels/workbook.xml.rels'].decode('utf-8')
    rel = re.search(r'<Relationship[^>]*Id="%s"[^>]*Target="([^"]+)"' % match.group(1), rels)
    if not rel:
        rel = re.search(r'<Relationship[^>]*Target="([^"]+)"[^>]*Id="%s"' % match.group(1), rels)
    target = rel.group(1).lstrip('/')
    return target if target.startswith('xl/') else 'xl/' + target


def patch_cells(path, sheet, transforms):
    """Apply regex transforms to a sheet's XML inside the saved workbook."""
    with zipfile.ZipFile(path) as archive:
        items = {item.filename: archive.read(item.filename) for item in archive.infolist()}
    target = sheet_xml_path(items, sheet)
    xml = items[target].decode('utf-8')
    for pattern, replacement in transforms:
        xml = re.sub(pattern, replacement, xml)
    items[target] = xml.encode('utf-8')
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, data in items.items():
            archive.writestr(name, data)


def add_cached_value(path, sheet, coord, value, is_error=False):
    """Give a saved formula cell a numeric cached <v> result."""
    t_attr = ' t="e"' if is_error else ''
    patch_cells(path, sheet, [
        (r'(<c r="%s"[^>]*)><f[^>]*>(.*?)</f>(?:(?!</c>).)*?</c>' % coord,
         r'\1%s><f>\2</f><v>%s</v></c>' % (t_attr, value))])


def make_shared_formula(path, sheet, master, follower, formula, master_value,
                        follower_value=None):
    """Convert two saved formula cells into a real shared-formula pair."""
    p_master = r'(<c r="%s"[^>]*)>(?:(?!</c>).)*?</c>' % master
    r_master = r'\1><f t="shared" ref="%s:%s" si="0">%s</f><v>%s</v></c>' % (master, follower, formula, master_value)
    p_foll = r'(<c r="%s"[^>]*)>(?:(?!</c>).)*?</c>' % follower
    v_foll = f'<v>{follower_value}</v>' if follower_value is not None else ''
    r_foll = r'\1><f t="shared" si="0"/>%s</c>' % v_foll
    patch_cells(path, sheet, [(p_master, r_master), (p_foll, r_foll)])


class ExtractorTestCase(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix='tpo-extract-test-'))
        self.addCleanup(shutil.rmtree, self.tmp, True)

    def source(self, name='source.xlsx', **kwargs):
        return build_source(self.tmp / name, **kwargs)

    def extract(self, **kwargs):
        return extractor.extract_monthly_inputs(self.source(**kwargs))

    def assert_error(self, pattern, **kwargs):
        with self.assertRaises(extractor.MonthlyInputError) as caught:
            self.extract(**kwargs)
        self.assertRegex(str(caught.exception), pattern)
        return str(caught.exception)


class TestExtraction(ExtractorTestCase):
    def test_extracts_prior_and_current_months_with_zero_actuals(self):
        pl = {2: dict(default_pl(2), revenueCents=0, cogsCents=0, sgaCents=0,
                      grossProfitCents=0, ebitCents=0, netIncomeCents=0)}
        pack = self.extract(pl=pl, sentinels_current=[
            True, True, False, False, False, False, False, False, False,
            False, False, False], report_month=2)
        self.assertEqual(pack['schema'], 'tpo-monthly-actuals-v1')
        self.assertEqual(pack['source']['reportMonth'], '2027-02')
        self.assertEqual(len(pack['months']), 14)
        self.assertEqual([m['month'] for m in pack['months']],
                         [f'2026-{m:02d}' for m in range(1, 13)] +
                         ['2027-01', '2027-02'])
        february = pack['months'][-1]
        self.assertEqual(february['financials'],
                         {'revenueCents': 0, 'cogsCents': 0, 'sgaCents': 0,
                          'netIncomeCents': 0})
        january = pack['months'][0]
        self.assertEqual(january['sourceRefs']['revenueCents'], ['PL!AG15'])
        self.assertEqual(january['sourceRefs']['arCents'], ['BS!Q9'])
        self.assertEqual(february['sourceRefs']['revenueCents'], ['PL!S15'])
        self.assertEqual(february['sourceRefs']['cashCents'],
                         ['BS!E5', 'BS!E6', 'BS!E7', 'BS!E8'])
        self.assertEqual(len(february['sourceRefs']['apCents']), 17)

    def test_misleading_header_years_are_ignored(self):
        pack = self.extract()
        self.assertEqual(pack['months'][0]['month'], '2026-01')
        self.assertEqual(pack['months'][-1]['month'], '2027-03')

    def test_negative_cogs_and_amounts_are_preserved(self):
        pl = {3: dict(default_pl(3), cogsCents=-125.5,
                      grossProfitCents=default_pl(3)['revenueCents'] + 125.5,
                      ebitCents=default_pl(3)['revenueCents'] + 125.5 - default_pl(3)['sgaCents'])}
        pack = self.extract(pl=pl)
        march = [m for m in pack['months'] if m['month'] == '2027-03'][0]
        self.assertEqual(march['financials']['cogsCents'], -12550)

    def test_half_up_cent_rounding(self):
        pl = {1: dict(default_pl(1), revenueCents=1.005,
                      grossProfitCents=1.005 - default_pl(1)['cogsCents'],
                      ebitCents=1.005 - default_pl(1)['cogsCents'] - default_pl(1)['sgaCents'])}
        bs = {1: dict(default_bs(1), cash=[1.004, 2.005, 0, 0])}
        pack = self.extract(pl=pl, bs=bs)
        january = [m for m in pack['months'] if m['month'] == '2027-01'][0]
        self.assertEqual(january['financials']['revenueCents'], 101)
        self.assertEqual(january['workingCapital']['cashCents'], 301)

    def test_formula_sentinel_genuine_zero_is_reported(self):
        path = self.source(sentinels_current=[True, True, True] + [False] * 9)
        wb = openpyxl.load_workbook(path)
        wb['PL']['T107'] = 0
        wb.save(path)
        pack = extractor.extract_monthly_inputs(path)
        self.assertEqual(pack['months'][-1]['month'], '2027-03')

    def test_borrowing_and_subtotal_rise_together_leave_ap_unchanged(self):
        month = 2
        plain = default_bs(month)
        inflated = default_bs(month)
        inflated['borrowing'] = [plain['borrowing'][0] + 9999.99,
                                 plain['borrowing'][1] + 8888.88]
        base = self.extract()
        changed = self.extract(bs={month: inflated})
        key = '2027-02'
        self.assertEqual(
            [m for m in changed['months'] if m['month'] == key][0]['workingCapital']['apCents'],
            [m for m in base['months'] if m['month'] == key][0]['workingCapital']['apCents'])

    def test_included_liability_change_moves_ap(self):
        month = 2
        tweaked = default_bs(month)
        tweaked['ap'] = list(tweaked['ap'])
        tweaked['ap'][1] += 100
        pack = self.extract(bs={month: tweaked})
        target = [m for m in pack['months'] if m['month'] == '2027-02'][0]
        self.assertGreater(target['workingCapital']['apCents'], 0)
        base = self.extract()
        baseline = [m for m in base['months'] if m['month'] == '2027-02'][0]
        self.assertEqual(target['workingCapital']['apCents'] -
                         baseline['workingCapital']['apCents'], 10000)

    def test_ap_detail_blank_is_zero_but_required_totals_are_errors(self):
        month = 1
        blank_detail = default_bs(month)
        blank_detail['ap'] = [None] + list(blank_detail['ap'][1:])
        pack = self.extract(bs={month: blank_detail})
        target = [m for m in pack['months'] if m['month'] == '2026-01'][0]
        base = self.extract()
        baseline = [m for m in base['months'] if m['month'] == '2026-01'][0]
        self.assertEqual(target['workingCapital']['apCents'],
                         baseline['workingCapital']['apCents'] - 1001)
        self.assert_error(r'BS!Q68: required amount is blank',
                          overrides=[('BS', 'Q68', None)])
        self.assert_error(r'BS!Q5: required amount is blank; blank is not zero',
                          overrides=[('BS', 'Q5', None)])


class TestLayoutGuards(ExtractorTestCase):
    def test_populated_sentinel_after_gap_is_rejected(self):
        self.assert_error(r'PL!U107: populated salary sentinel after a gap',
                          sentinels_current=[True, True, False, True, False,
                                             False, False, False, False, False,
                                             False, False])

    def test_run_must_end_exactly_at_title_month(self):
        self.assert_error(r'actual run ends at month 2, but the title reports 3',
                          sentinels_current=[True, True, False, False, False,
                                             False, False, False, False, False,
                                             False, False])
        self.assert_error(r'actual run ends at month 4, but the title reports 3',
                          sentinels_current=[True] * 4 + [False] * 8)

    def test_incomplete_prior_year_is_rejected(self):
        self.assert_error(r'prior comparison year 2026 has only 11/12',
                          sentinels_prior=[True] * 11 + [False])

    def test_disagreeing_titles_are_rejected(self):
        self.assert_error(r'PL!A2 / BS!A2: report titles disagree',
                          titles={'BS': 'Balance Sheet  :  YTD February 2027'})

    def test_year_and_month_anchors_are_enforced(self):
        self.assert_error(r'PL!AD3', titles={'AD3': 'YTD-26'})
        self.assert_error(r'PL!AS3', titles={'AS3': 'YTD-25'})
        self.assert_error(r'IncomeStatement!R6',
                          overrides=[('IncomeStatement', 'R6', 2025)])
        self.assert_error(r'IncomeStatement!AS7',
                          overrides=[('IncomeStatement', 'AS7', 'Jan-27')])
    def test_year_anchor_fractional_is_rejected(self):
        self.assert_error(r'IncomeStatement!R6',
                          overrides=[('IncomeStatement', 'R6', 2027.9)])

    def test_bs_inventory_label_without_product_or_material_is_rejected(self):
        self.assert_error(r'BS!B14: inventory-component label does not match',
                          overrides=[('BS', 'B14', 'Packing')])
    def test_shifted_account_codes_and_labels_are_rejected(self):
        self.assert_error(r'PL!A63: account anchor',
                          overrides=[('PL', 'A63', 'T-X')])
        self.assert_error(r'PL!B65: anchor label',
                          overrides=[('PL', 'B65', 'Gross Marginal Profit')])
        self.assert_error(r'BS!A49: account anchor',
                          overrides=[('BS', 'A49', 21398)])
        self.assert_error(r'BS!B68: expected the "Total Liabilities"',
                          overrides=[('BS', 'B68', 'All Liabilities')])
        self.assert_error(r'BS!B53: expected a borrowing label',
                          overrides=[('BS', 'B53', 'Loan from related party')])
        self.assert_error(r'BS!B9: expected a trade-receivables label',
                          overrides=[('BS', 'B9', 'Other receivables')])

    def test_missing_sheet_reports_name(self):
        path = self.source()
        wb = openpyxl.load_workbook(path)
        wb.remove(wb['IncomeStatement'])
        wb.save(path)
        with self.assertRaises(extractor.MonthlyInputError) as caught:
            extractor.extract_monthly_inputs(path)
        self.assertIn('required sheet "IncomeStatement" is missing',
                      str(caught.exception))


class TestFormulaCaches(ExtractorTestCase):
    def test_formula_without_cache_is_an_error(self):
        path = self.source(overrides=[('PL', 'S15', '=SUM(S6:S14)')])
        with self.assertRaises(extractor.MonthlyInputError) as caught:
            extractor.extract_monthly_inputs(path)
        self.assertIn('PL!S15: formula has no numeric cached value',
                      str(caught.exception))

    def test_formula_with_cache_is_used(self):
        path = self.source(overrides=[
            ('PL', 'R15', '=SUBTOTAL(9,R6:R14)'),
            ('PL', 'R65', 3817.25),
            ('PL', 'R172', 3414.5),
            ('PL', 'AG15', '=SUBTOTAL(9,AG6:AG14)'),
            ('PL', 'AG65', 8260.75),
            ('PL', 'AG172', 7858.0)])
        add_cached_value(path, 'PL', 'R15', 4321.75)
        add_cached_value(path, 'PL', 'AG15', 8765.25)
        pack = extractor.extract_monthly_inputs(path)
        january = [m for m in pack['months'] if m['month'] == '2027-01'][0]
        prior = [m for m in pack['months'] if m['month'] == '2026-01'][0]
        self.assertEqual(january['financials']['revenueCents'], 432175)
        self.assertEqual(prior['financials']['revenueCents'], 876525)

    def test_shared_formula_follower_without_cache_is_an_error(self):
        path = self.source(overrides=[('PL', 'R15', '=SUM(R6:R14)'),
                                      ('PL', 'R65', 495.5),
                                      ('PL', 'R172', 92.75),
                                      ('PL', 'S15', '=SUM(S6:S14)')])
        make_shared_formula(path, 'PL', 'R15', 'S15', 'SUM(R6:R14)', 1000)
        with self.assertRaises(extractor.MonthlyInputError) as caught:
            extractor.extract_monthly_inputs(path)
        self.assertIn('PL!S15: formula has no numeric cached value',
                      str(caught.exception))
    def test_shared_formula_follower_with_cache_is_used(self):
        path = self.source(overrides=[
            ('PL', 'R15', '=SUM(R6:R14)'),
            ('PL', 'R65', 730.0),
            ('PL', 'R172', 327.25),
            ('PL', 'S15', '=SUM(S6:S14)'),
            ('PL', 'S65', 1837.25),
            ('PL', 'S172', 1432.5)])
        make_shared_formula(path, 'PL', 'R15', 'S15', 'SUM(R6:R14)', 1234.5,
                            follower_value=2345.75)
        pack = extractor.extract_monthly_inputs(path)
        february = [m for m in pack['months'] if m['month'] == '2027-02'][0]
        self.assertEqual(february['financials']['revenueCents'], 234575)

    def test_formula_sentinel_missing_cache_is_an_error(self):
        path = self.source(overrides=[('PL', 'R107', '=SUM(R100:R106)')])
        with self.assertRaises(extractor.MonthlyInputError) as caught:
            extractor.extract_monthly_inputs(path)
        self.assertIn('PL!R107: salary-sentinel formula has no numeric cached value',
                      str(caught.exception))

    def test_formula_sentinel_error_cache_is_an_error(self):
        path = self.source(overrides=[('PL', 'R107', '=SUM(R100:R106)')])
        add_cached_value(path, 'PL', 'R107', '#VALUE!', is_error=True)
        with self.assertRaises(extractor.MonthlyInputError) as caught:
            extractor.extract_monthly_inputs(path)
        self.assertIn('PL!R107: salary-sentinel formula has no numeric cached value',
                      str(caught.exception))

    def test_ap_formula_cache_missing_distinct_from_absent_leaf(self):
        path = self.source(overrides=[('BS', 'Q48', '=SUM(Q30:Q40)')])
        with self.assertRaises(extractor.MonthlyInputError) as caught:
            extractor.extract_monthly_inputs(path)
        self.assertIn('BS!Q48: formula has no numeric cached value',
                      str(caught.exception))

    def test_boolean_in_required_amount_is_rejected(self):
        path = self.source(overrides=[('PL', 'R15', True)])
        with self.assertRaises(extractor.MonthlyInputError) as caught:
            extractor.extract_monthly_inputs(path)
        self.assertIn('PL!R15: boolean value is not a valid amount',
                      str(caught.exception))

    def test_non_finite_number_is_rejected(self):
        class DummyView:
            def __getitem__(self, sheet):
                return {'R15': type('Cell', (), {'value': float('nan')})()}
        with self.assertRaises(extractor.MonthlyInputError) as caught:
            extractor._amount(DummyView(), DummyView(), 'PL', 15, 18)
        self.assertIn('PL!R15: non-finite numeric value',
                      str(caught.exception))

    def test_formula_error_cache_is_rejected(self):
        path = self.source(overrides=[('PL', 'R15', '=SUM(R6:R14)')])
        add_cached_value(path, 'PL', 'R15', '#DIV/0!', is_error=True)
        with self.assertRaises(extractor.MonthlyInputError) as caught:
            extractor.extract_monthly_inputs(path)
        self.assertIn('PL!R15: cached value', str(caught.exception))

class TestReviewWorkbook(ExtractorTestCase):
    def test_review_roundtrip_and_source_immutability(self):
        source = self.source(name='=TPO_Review_Test_2026.xlsx')
        digest_before = source.read_bytes()
        pack = extractor.extract_monthly_inputs(source)
        output = self.tmp / 'review.xlsx'
        extractor.write_review_workbook(pack, output)
        self.assertEqual(source.read_bytes(), digest_before)
        review = openpyxl.load_workbook(output)
        self.assertEqual(review.sheetnames,
                         ['MonthlyFinancials', '1. Working Capital', 'Source Map'])
        fin = list(review['MonthlyFinancials'].iter_rows(values_only=True))
        self.assertEqual(fin[0][:8], ('Month', 'Total Revenue', 'COGS',
                                      'Gross Profit', 'SG&A', 'EBIT', 'Net Income',
                                      'Quarter'))
        self.assertEqual(len(fin) - 1, len(pack['months']))
        for row, record in zip(fin[1:], pack['months']):
            self.assertEqual(row[0], extractor._month_label(record['month']))
            f = record['financials']
            gp = f['revenueCents'] - f['cogsCents']
            ebit = gp - f['sgaCents']
            for col, cents in zip(range(1, 7),
                                  [f['revenueCents'],
                                   f['cogsCents'],
                                   gp,
                                   f['sgaCents'],
                                   ebit,
                                   f['netIncomeCents']]):
                self.assertEqual(round(row[col] * 100), cents)
        wc = list(review['1. Working Capital'].iter_rows(values_only=True))
        for row, record in zip(wc[1:], pack['months']):
            w = record['workingCapital']
            for col, cents in zip(range(1, 6),
                                  [w['cashCents'], w['arCents'], w['inventoryCents'],
                                   w['apCents'],
                                   w['cashCents'] + w['arCents'] + w['inventoryCents'] - w['apCents']]):
                self.assertEqual(round(row[col] * 100), cents)
        source_map = review['Source Map']
        self.assertIn('Review copy only', source_map['A1'].value)
        self.assertEqual(source_map['B3'].value, pack['source']['filename'])
        self.assertEqual(source_map['B3'].data_type, 's')
        self.assertEqual(source_map['B4'].value, pack['source']['sha256'])
        rows = list(source_map.iter_rows(values_only=True))
        self.assertTrue(any(r[1] == 'Borrowing excluded from Accounts Payable'
                            and 'BS!S53' in str(r[2]) for r in rows))
        for sheet in review.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    self.assertNotEqual(cell.data_type, 'f', f'{sheet.title}!{cell.coordinate} has formula data_type')


class TestCommandLine(ExtractorTestCase):
    def run_main(self, *args):
        return extractor.main(list(args))

    def test_happy_path_writes_both_outputs(self):
        source = self.source()
        prefix = self.tmp / 'out/pack'
        (self.tmp / 'out').mkdir()
        self.assertEqual(self.run_main('--source', str(source),
                                       '--output-prefix', str(prefix)), 0)
        pack = json.loads(Path(f'{prefix}.json').read_text())
        self.assertEqual(len(pack['months']), 15)
        self.assertTrue(Path(f'{prefix}.xlsx').exists())

    def test_refuses_to_overwrite_existing_outputs(self):
        source = self.source()
        prefix = self.tmp / 'pack'
        Path(f'{prefix}.json').write_text('old')
        self.assertEqual(self.run_main('--source', str(source),
                                       '--output-prefix', str(prefix)), 2)
        self.assertEqual(Path(f'{prefix}.json').read_text(), 'old')
        self.assertFalse(Path(f'{prefix}.xlsx').exists())

    def test_refuses_output_resolving_to_source(self):
        source = self.source(name='book.xlsx')
        self.assertEqual(self.run_main('--source', str(source),
                                       '--output-prefix',
                                       str(source.with_suffix(''))), 2)

    def test_invalid_input_creates_no_outputs(self):
        source = self.source(sentinels_current=[True, True, False, True] + [False] * 8)
        prefix = self.tmp / 'none'
        self.assertEqual(self.run_main('--source', str(source),
                                       '--output-prefix', str(prefix)), 2)
        self.assertFalse(Path(f'{prefix}.json').exists())
        self.assertFalse(Path(f'{prefix}.xlsx').exists())

    def test_missing_source_fails(self):
        self.assertEqual(self.run_main('--source', str(self.tmp / 'nope.xlsx'),
                                       '--output-prefix',
                                       str(self.tmp / 'x')), 2)


if __name__ == '__main__':
    unittest.main()
